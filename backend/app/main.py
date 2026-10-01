import csv, io, json, os, shutil, sqlite3, tempfile, uuid
from datetime import date, datetime
from fastapi import BackgroundTasks, FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel
from .cleaner import CLEANERS, clean_row, s
from .config import FIELDS, TENANTS
from .dedupe import find_links, outliers
from . import memory

DB = os.getenv("DB_PATH", "loan.db"); app = FastAPI(title="Loan Register Cleaner")

def db():
    c = sqlite3.connect(DB); c.row_factory = sqlite3.Row; return c

with db() as c:
    c.executescript("""CREATE TABLE IF NOT EXISTS jobs(id TEXT PRIMARY KEY,tenant TEXT,filename TEXT,status TEXT,error TEXT);
    CREATE TABLE IF NOT EXISTS rows(job TEXT,n INT,raw TEXT,clean TEXT,fixes TEXT,flags TEXT,info TEXT,state TEXT,PRIMARY KEY(job,n));
    CREATE INDEX IF NOT EXISTS ix ON rows(job,state);
    CREATE TABLE IF NOT EXISTS decisions(id INTEGER PRIMARY KEY,tenant TEXT,rule TEXT,ctx TEXT,vec TEXT,action TEXT,value TEXT);""")

def stream_records(path, t):
    """Yield dicts keyed by canonical field names without loading the whole file (xlsx read_only / csv)."""
    if path.lower().endswith(".csv"):
        it = csv.reader(open(path, newline="", encoding="utf-8-sig"))
    else:
        it = load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True)
    hdr = [s(h).lower() for h in next(it)]
    idx = {k: next((i for i, h in enumerate(hdr) if h in al), -1) for k, al in t.cols.items()}
    miss = [f for f in FIELDS if idx[f] < 0]
    if miss: raise ValueError(f"Missing column(s): {', '.join(miss)}. Add header aliases in TENANTS config.")
    for r in it:
        if any(s(c) for c in r): yield {k: (r[i] if i >= 0 and i < len(r) else "") for k, i in idx.items()}

def state_of(flags, excluded=False):
    return "excluded" if excluded else ("review" if any(not f["resolved"] for f in flags) else "ok")

def process(job, path, tenant):
    c, t = db(), TENANTS[tenant]
    try:
        index, batch = [], []
        for n, raw in enumerate(stream_records(path, t), start=2):           # pass 1: stream + clean
            clean, fixes, flags = clean_row(raw, t)
            index.append((n, clean["name"], clean["amount"], clean["date"], clean["phone"], any(f["field"] == "date" for f in flags)))
            raw = {k: (v.isoformat() if isinstance(v, (date, datetime)) else v) for k, v in raw.items()}
            batch.append((job, n, json.dumps(raw, default=str), json.dumps(clean), json.dumps(fixes), json.dumps(flags), "[]", ""))
            if len(batch) >= 5000: c.executemany("INSERT INTO rows VALUES(?,?,?,?,?,?,?,?)", batch); batch = []
        c.executemany("INSERT INTO rows VALUES(?,?,?,?,?,?,?,?)", batch)
        lf, li = find_links(index, t); ol = outliers(index, t)                # pass 2: indexed cross-row checks
        for n in {*lf, *li, *ol}:
            r = c.execute("SELECT flags,info FROM rows WHERE job=? AND n=?", (job, n)).fetchone()
            flags = json.loads(r["flags"]) + lf.get(n, []) + ([ol[n]] if n in ol else [])
            c.execute("UPDATE rows SET flags=?,info=? WHERE job=? AND n=?", (json.dumps(flags), json.dumps(li.get(n, [])), job, n))
        for n, fl in c.execute("SELECT n,flags FROM rows WHERE job=?", (job,)).fetchall():
            c.execute("UPDATE rows SET state=? WHERE job=? AND n=?", (state_of(json.loads(fl)), job, n))
        c.execute("UPDATE jobs SET status='done' WHERE id=?", (job,))
    except Exception as e:
        c.execute("UPDATE jobs SET status='error',error=? WHERE id=?", (str(e), job))
    c.commit(); os.unlink(path)

@app.post("/api/jobs")
async def create_job(bg: BackgroundTasks, file: UploadFile = File(...), tenant: str = "default"):
    if tenant not in TENANTS: raise HTTPException(400, "Unknown tenant")
    job = uuid.uuid4().hex[:10]; ext = os.path.splitext(file.filename)[1]
    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as f: shutil.copyfileobj(file.file, f)
    with db() as c: c.execute("INSERT INTO jobs VALUES(?,?,?,?,NULL)", (job, tenant, file.filename, "processing"))
    bg.add_task(process, job, f.name, tenant)       # swap for Celery/RQ when files get big
    return {"id": job}

@app.get("/api/jobs/{job}")
def job_status(job: str):
    c = db(); j = c.execute("SELECT * FROM jobs WHERE id=?", (job,)).fetchone()
    if not j: raise HTTPException(404)
    cnt = {r["state"]: r["k"] for r in c.execute("SELECT state,COUNT(*) k FROM rows WHERE job=? GROUP BY state", (job,))}
    return {**dict(j), "counts": cnt, "fixes": c.execute("SELECT SUM(json_array_length(fixes)) FROM rows WHERE job=?", (job,)).fetchone()[0] or 0}

@app.get("/api/jobs/{job}/rows")
def rows(job: str, state: str = "review", page: int = 1, size: int = 25):
    c = db(); j = c.execute("SELECT tenant FROM jobs WHERE id=?", (job,)).fetchone()
    q, a = ("SELECT * FROM rows WHERE job=?", [job]) if state == "all" else ("SELECT * FROM rows WHERE job=? AND state=?", [job, state])
    out = []
    for r in c.execute(q + " ORDER BY n LIMIT ? OFFSET ?", a + [size, (page - 1) * size]):
        flags = json.loads(r["flags"])
        for f in flags:
            if not f["resolved"]: f["suggestion"] = memory.suggest(c, j["tenant"], f)       # RAG suggestion per flag
        out.append({"n": r["n"], "raw": json.loads(r["raw"]), "clean": json.loads(r["clean"]), "fixes": json.loads(r["fixes"]),
                    "flags": flags, "info": json.loads(r["info"]), "state": r["state"]})
    return out

class Resolve(BaseModel):
    action: str          # apply | keep | blank | exclude
    value: str | None = None

@app.post("/api/jobs/{job}/rows/{n}/flags/{i}")
def resolve(job: str, n: int, i: int, body: Resolve):
    c = db(); j = c.execute("SELECT tenant FROM jobs WHERE id=?", (job,)).fetchone()
    r = c.execute("SELECT * FROM rows WHERE job=? AND n=?", (job, n)).fetchone()
    clean, flags, fixes = json.loads(r["clean"]), json.loads(r["flags"]), json.loads(r["fixes"]); f = flags[i]
    if body.action == "apply":                       # re-validate with the same rule -> a bad value can't clear a flag
        res = CLEANERS[f["field"]](body.value, TENANTS[j["tenant"]])
        if "err" in res and res["kind"] != "confirm": raise HTTPException(422, res["err"])
        fixes.append({"field": f["field"], "from": s(json.loads(r["raw"]).get(f["field"])), "to": res["v"], "msg": "Corrected by reviewer", "by": "user"})
        clean[f["field"]] = res["v"]
    elif body.action == "blank": clean[f["field"]] = ""; fixes.append({"field": f["field"], "from": "", "to": "", "msg": "Left blank by reviewer", "by": "user"})
    f["resolved"] = True; st = state_of(flags, body.action == "exclude")
    c.execute("UPDATE rows SET clean=?,flags=?,fixes=?,state=? WHERE job=? AND n=?", (json.dumps(clean), json.dumps(flags), json.dumps(fixes), st, job, n))
    memory.remember(c, j["tenant"], f, body.action, body.value)      # every click trains the suggestions
    c.commit(); return {"state": st}

@app.get("/api/jobs/{job}/export")
def export(job: str):
    c = db(); wb = Workbook(write_only=True); T = ["branch", "customer_name", "loan_amount", "loan_date", "phone", "status"]
    sh = {k: wb.create_sheet(k) for k in ["Clean", "Needs Review", "Excluded", "Change Log"]}
    sh["Clean"].append(T); sh["Needs Review"].append(T + ["source_row", "why_flagged"])
    sh["Excluded"].append(["source_row", "customer_name", "reason"]); sh["Change Log"].append(["source_row", "field", "before", "after", "what_changed", "by"])
    for r in c.execute("SELECT * FROM rows WHERE job=? ORDER BY n", (job,)):
        cl, fl = json.loads(r["clean"]), json.loads(r["flags"]); g = [cl[f] for f in FIELDS]
        if r["state"] == "ok": sh["Clean"].append(g)
        elif r["state"] == "review": sh["Needs Review"].append(g + [r["n"], " | ".join(f["msg"] for f in fl if not f["resolved"])])
        else: sh["Excluded"].append([r["n"], cl["name"], " | ".join(f["msg"] for f in fl)])
        for x in json.loads(r["fixes"]): sh["Change Log"].append([r["n"], x["field"], x["from"], x["to"], x["msg"], x["by"]])
    p = tempfile.mktemp(suffix=".xlsx"); wb.save(p)
    return FileResponse(p, filename="Clean_Loan_Register.xlsx")
